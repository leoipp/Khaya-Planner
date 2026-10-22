import io
import unittest
import warnings
import pandas as pd
from openpyxl import load_workbook
from app.screens.simulacao import _escrever_paginas_excel, _particionar_paginas_excel


class ExportacaoPaginasExcel(unittest.TestCase):
    def test_preserva_linhas_colunas_vazias_e_schemas_diferentes_sem_aviso(self):
        a = pd.DataFrame({'cenario':['Sem desbaste'], 'idade':[None], 'volume':[10.5], 'vazia':[None]})
        b = pd.DataFrame({'cenario':['Com desbaste'], 'idade':[6.], 'volume':[20.], 'vazia':[None], 'extra':[3.]})
        vazio = pd.DataFrame(columns=['so_no_vazio'])
        paginas, excedidos = _particionar_paginas_excel([('A', a), ('B', b), ('Vazio', vazio)])
        self.assertEqual(excedidos, [])
        arquivo = io.BytesIO()
        with warnings.catch_warnings():
            warnings.simplefilter('error', FutureWarning)
            with pd.ExcelWriter(arquivo, engine='xlsxwriter') as writer:
                _escrever_paginas_excel(writer, paginas, 'Simulação')
        arquivo.seek(0)
        livro = load_workbook(arquivo, read_only=True)
        try:
            self.assertEqual(list(livro['Simulação'].values), [
                ('cenario', 'idade', 'volume', 'vazia', 'extra', 'so_no_vazio'),
                ('Sem desbaste', None, 10.5, None, None, None),
                ('Com desbaste', 6, 20, None, 3, None)])
        finally:livro.close()

    def test_paginas_distintas_preservam_cenarios_inteiros(self):
        a = pd.DataFrame({'cenario':['A', 'A'], 'volume':[1., 2.]})
        b = pd.DataFrame({'cenario':['B', 'B'], 'volume':[3., 4.]})
        paginas, excedidos = _particionar_paginas_excel([('A', a), ('B', b)], limite_linhas=3)
        self.assertEqual(len(paginas), 2)
        self.assertEqual(excedidos, [])
        arquivo = io.BytesIO()
        with pd.ExcelWriter(arquivo, engine='xlsxwriter') as writer:
            _escrever_paginas_excel(writer, paginas, 'Simulação')
        arquivo.seek(0)
        livro = load_workbook(arquivo, read_only=True)
        try:
            self.assertEqual(livro.sheetnames, ['Simulação 1', 'Simulação 2'])
            self.assertEqual(list(livro['Simulação 2'].values), [('cenario', 'volume'), ('B', 3), ('B', 4)])
        finally:livro.close()
