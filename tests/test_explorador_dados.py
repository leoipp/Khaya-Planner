import os,sqlite3,tempfile,time,unittest
from pathlib import Path
from unittest.mock import patch
from contextlib import closing
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from openpyxl import load_workbook
from app.core import explorador_dados as d


class ExploradorCore(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.banco=str(Path(self.tmp.name,'banco.sqlite'))
        with closing(sqlite3.connect(self.banco)) as c:
            c.execute('CREATE TABLE dados(id INTEGER, grupo TEXT, valor REAL, texto TEXT, arquivo BLOB)')
            c.executemany('INSERT INTO dados VALUES(?,?,?,?,?)',[(i,'A' if i%2==0 else 'B',float(i),'=1+1',b'abc') for i in range(1005)])
            c.execute('CREATE VIEW vista AS SELECT id,valor FROM dados');c.commit()
    def test_paginacao_consulta_cte_view_e_blob(self):
        self.assertEqual(d.tabelas(self.banco),['dados','vista'])
        r=d.pagina(self.banco,'WITH t AS (SELECT * FROM dados) SELECT * FROM t ORDER BY id',3,500)
        self.assertEqual(len(r['linhas']),5);self.assertFalse(r['mais']);self.assertEqual(r['linhas'][0][0],1000)
        self.assertEqual(r['linhas'][0][-1],'[BLOB 3 bytes]')
        self.assertTrue(d.pagina(self.banco,'SELECT * FROM vista',1,100)['mais'])
    def test_consulta_nao_modifica_banco(self):
        for sql in ['DELETE FROM dados','SELECT 1; DROP TABLE dados','ATTACH DATABASE ":memory:" AS outro',"SELECT load_extension('inexistente')"]:
            with self.assertRaises((sqlite3.DatabaseError,ValueError)):d.pagina(self.banco,sql)
        self.assertEqual(len(d.pagina(self.banco,'SELECT COUNT(*) FROM dados')['linhas']),1)
        self.assertEqual(d.pagina(self.banco,'SELECT COUNT(*) FROM dados')['linhas'][0][0],1005)
    def test_excel_exporta_consulta_inteira_em_varias_abas(self):
        destino=Path(self.tmp.name,'dados.xlsx');r=d.exportar_excel(self.banco,'SELECT id,texto FROM dados ORDER BY id',destino,linhas_por_aba=500)
        self.assertEqual(r['linhas'],1005);self.assertEqual(r['abas'],3)
        livro=load_workbook(destino,read_only=True)
        try:
            self.assertEqual(livro.sheetnames,['Dados_1','Dados_2','Dados_3'])
            self.assertEqual(list(livro['Dados_3'].values)[-1],(1004,'=1+1'))
            self.assertEqual(livro['Dados_1']['B2'].data_type,'s')
        finally:livro.close()
    def test_grafico_usa_todas_linhas_e_valida_coluna_numerica(self):
        r=d.dados_grafico(self.banco,'SELECT * FROM dados','grupo','valor','sum')
        self.assertEqual(r,[('A',float(sum(range(0,1005,2)))),('B',float(sum(range(1,1005,2))))])
        self.assertEqual(d.dados_grafico(self.banco,'SELECT * FROM dados','grupo','','count'),[('A',503),('B',502)])
        with self.assertRaises(ValueError):d.dados_grafico(self.banco,'SELECT * FROM dados','id','texto')
        with self.assertRaises(ValueError):d.dados_grafico(self.banco,'SELECT * FROM dados','id','valor','raw',limite=10)
    def test_exportacao_com_erro_preserva_destino_anterior(self):
        destino=Path(self.tmp.name,'dados.xlsx');destino.write_bytes(b'anterior')
        with self.assertRaises(sqlite3.DatabaseError):d.exportar_excel(self.banco,'SELECT inexistente FROM dados',destino)
        self.assertEqual(destino.read_bytes(),b'anterior');self.assertEqual(list(Path(self.tmp.name).glob('*.tmp.xlsx')),[])


class ExploradorUI(unittest.TestCase):
    def setUp(self):ExploradorCore.setUp(self)
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app=QApplication.instance() or QApplication([])
    def aguardar(self,tela):
        prazo=time.monotonic()+15
        while tela.ocupado and time.monotonic()<prazo:self.app.processEvents();time.sleep(.01)
        self.assertFalse(tela.ocupado);self.assertFalse(tela.status.text().startswith('Erro:'),tela.status.text())
    def test_tela_pagina_ordena_exporta_e_grafica(self):
        from app.screens.explorador_dados import TelaExploradorDados
        with patch('app.screens.explorador_dados.projeto.projeto_aberto',return_value=False):tela=TelaExploradorDados()
        try:
            tela.abrir_tabela(self.banco,'dados');self.aguardar(tela)
            self.assertEqual(tela.modelo.rowCount(),500);self.assertTrue(tela.proxima.isEnabled())
            tela._pagina(3);self.aguardar(tela);self.assertEqual(tela.modelo.linhas[0][0],1000)
            tela._ordenar(0);self.aguardar(tela);tela._ordenar(0);self.aguardar(tela)
            self.assertEqual(tela.modelo.linhas[0][0],1004)
            tela.tamanho.setCurrentIndex(0);self.aguardar(tela);self.assertEqual(tela._pagina_atual,1);self.assertEqual(tela.modelo.rowCount(),100)
            destino=str(Path(self.tmp.name,'tela.xlsx'))
            with patch('app.screens.explorador_dados.QFileDialog.getSaveFileName',return_value=(destino,'')):tela._exportar()
            self.aguardar(tela);self.assertIn('1,005 linhas',tela.status.text())
            tela.eixo_x.setCurrentText('grupo');tela.eixo_y.setCurrentText('valor');tela._grafico();self.aguardar(tela)
            self.assertIsNotNone(tela._figura);self.assertEqual(len(tela._figura.axes[0].patches),2)
        finally:tela.close();tela.deleteLater();self.app.processEvents()
